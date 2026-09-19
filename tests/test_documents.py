import io

from mailpulse.documents.extract import (
    PDF_TEXT_MIN_CHARS,
    extract_attachment,
    extract_docx_text,
    extract_html_text,
)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for row, cells in zip(t.rows, table, strict=True):
            for cell, value in zip(row.cells, cells, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class FakeVision:
    def __init__(self, text: str = "СЧЁТ №507 на 320 000 тенге, оплатить до 18.09.2026") -> None:
        self.text = text
        self.calls: list[tuple[str, str]] = []
        self.metas: list[dict] = []

    async def __call__(self, raw: bytes, mime: str, filename: str) -> str:
        self.calls.append((mime, filename))
        return self.text


def test_docx_extracts_paragraphs_and_tables():
    raw = make_docx(
        ["Договор аренды", "Ставка 6500 тенге"],
        table=[["Услуга", "Цена"], ["Уборка", "180000"]],
    )

    result = extract_docx_text(raw)

    assert "Договор аренды" in result
    assert "Услуга | Цена" in result
    assert "Уборка | 180000" in result


def test_html_strips_markup():
    raw = (
        "<html><body><p>Добавить <b>Add to cart</b></p><script>x=1</script></body></html>".encode()
    )

    result = extract_html_text(raw)

    assert "Add to cart" in result
    assert "script" not in result


async def test_docx_routed_by_extension():
    raw = make_docx(["Служебная записка"])

    extracted = await extract_attachment(
        raw, "application/octet-stream", "записка.docx", vision=None
    )

    assert extracted.method == "docx"
    assert "Служебная записка" in extracted.text


async def test_scanned_pdf_without_text_goes_to_vision():
    vision = FakeVision()

    extracted = await extract_attachment(blank_pdf(), "application/pdf", "scan.pdf", vision=vision)

    assert extracted.method == "vision"
    assert "320 000" in extracted.text
    assert vision.calls == [("application/pdf", "scan.pdf")]


async def test_image_goes_to_vision():
    vision = FakeVision()

    extracted = await extract_attachment(
        b"\xff\xd8fakejpeg", "image/jpeg", "photo.jpg", vision=vision
    )

    assert extracted.method == "vision"
    assert vision.calls == [("image/jpeg", "photo.jpg")]


async def test_image_without_vision_returns_none():
    extracted = await extract_attachment(b"\xff\xd8", "image/jpeg", "photo.jpg", vision=None)

    assert extracted == extracted.__class__("", "none")


async def test_vision_failure_degrades_to_none():
    async def boom(raw, mime, filename):
        raise RuntimeError("api down")

    extracted = await extract_attachment(b"\xff\xd8", "image/png", "x.png", vision=boom)

    assert extracted.method == "none"
    assert extracted.text == ""


async def test_blank_pdf_has_no_text_layer():
    from mailpulse.documents.extract import extract_pdf_text

    assert len(extract_pdf_text(blank_pdf())) < PDF_TEXT_MIN_CHARS


async def test_unknown_type_is_skipped():
    extracted = await extract_attachment(
        b"PK\x03\x04zip", "application/zip", "archive.zip", vision=FakeVision()
    )

    assert extracted.method == "none"
